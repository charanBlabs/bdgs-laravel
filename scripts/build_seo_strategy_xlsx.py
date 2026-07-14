"""Build docs/seo-keyword-strategy-planner.xlsx from workspace sources only.

Sources:
- data/seo-pages.json
- docs/seo-keyword-map.md (research status facts only — does not modify that file)
- Live solutions from DB / seed (scripts/_solutions_list.json via artisan helper)

Does NOT include Google HubSpoke / Keyword Mapping restructuring sheets.
"""

from __future__ import annotations

import json
import subprocess
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "seo-keyword-strategy-planner.xlsx"
SOLUTIONS_JSON = ROOT / "scripts" / "_solutions_list.json"
SEO_PAGES = ROOT / "data" / "seo-pages.json"

header_fill = PatternFill("solid", fgColor="1F4E79")
header_font = Font(bold=True, color="FFFFFF", size=11)
live_fill = PatternFill("solid", fgColor="D5F5E3")
planned_fill = PatternFill("solid", fgColor="FCF3CF")
research_done = PatternFill("solid", fgColor="ABEBC6")
research_partial = PatternFill("solid", fgColor="F9E79F")
research_needed = PatternFill("solid", fgColor="F5B7B1")
thin = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(wrap_text=True, vertical="top", horizontal="center")

# Keyword research status from docs/seo-keyword-map.md (workspace work so far)
RESEARCH = {
    "/": ("Partial — confirm", "brilliant directories developers", "directory website expert", "FAQ questions still empty"),
    "/setup": ("Partial — confirm", "brilliant directories setup", "brilliant directories setup cost", "Cost FAQ: how much does brilliant directories cost"),
    "/hire-developer": ("Partial — confirm", "hire brilliant directories developer", "dedicated bd developer", "Need secondaries + FAQs"),
    "/consultation": ("Partial — confirm", "brilliant directories consultant", "directory website expert (shared with /)", "Confirm intent split vs /"),
    "/reviews": (
        "Researched",
        "bd growth suite reviews",
        "verified client reviews; Marketplace reviews; Business Labs by BD Growth Suite; Gold Certified partner",
        "Ready for Yakin review/lock. ≠ platform 'brilliant directories reviews'",
    ),
    "/zoom-clinics": (
        "Researched",
        "brilliant directories zoom clinics",
        "bd growth suite zoom clinics; free Zoom Clinics; live Q&A",
        "Ready for Yakin review/lock. Do not primary 'brilliant directories help'",
    ),
    "/solutions/seo": (
        "Researched",
        "brilliant directories schema solutions",
        "SEO & Schema Solutions for Brilliant Directories; structured data; JSON-LD; Essential/Strategic SEO Audit",
        "Ready for Yakin review. Registry still has wrong primary 'brilliant directories schema'. Product long-tails → /solutions/[slug]",
    ),
}

CAT_META = {
    "seo": ("SEO & Schema", "/solutions/seo", "brilliant directories schema solutions"),
    "lead-gen": ("Lead Generation & Conversion", "/solutions/lead-gen", "directory lead generation"),
    "member-profiles": ("Member Profile Enhancement", "/solutions/member-profiles", "directory member profiles"),
    "search": ("Search & Discovery", "/solutions/search", "directory search filters"),
    "page-design": ("Page Design & Development", "/solutions/page-design", "directory page design"),
    "content": ("Content & Engagement", "/solutions/content", "directory content solutions"),
    "integrations": ("Integrations", "/solutions/integrations", "brilliant directories integrations"),
    "member-management": ("Member Management & Automation", "/solutions/member-management", "directory member management"),
    "": ("Uncategorized solutions", "/solutions", "— (assign category)"),
}


def style_header(ws, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(1, c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_rows(ws, headers, rows, status_col=None, research_col=None):
    for col, h in enumerate(headers, 1):
        ws.cell(1, col, h)
    style_header(ws, len(headers))
    for r_idx, row in enumerate(rows, 2):
        for c_idx, val in enumerate(row, 1):
            cell = ws.cell(r_idx, c_idx, val)
            cell.alignment = wrap
            cell.border = thin
            if status_col and c_idx == status_col:
                s = str(val or "")
                if s.lower() in ("live", "published"):
                    cell.fill = live_fill
                elif s.lower() in ("planned", "uncategorized"):
                    cell.fill = planned_fill
                cell.alignment = center
            if research_col and c_idx == research_col:
                s = str(val or "")
                if s == "Researched":
                    cell.fill = research_done
                elif "Partial" in s:
                    cell.fill = research_partial
                elif "Needs research" in s:
                    cell.fill = research_needed
                cell.alignment = center


def load_solutions() -> list[dict]:
    helper = ROOT / "scripts" / "_list_solutions.php"
    if helper.exists():
        subprocess.run(["php", str(helper)], cwd=ROOT, check=False, capture_output=True)
    if not SOLUTIONS_JSON.exists():
        return []
    return json.loads(SOLUTIONS_JSON.read_text(encoding="utf-8"))


def load_seo_pages() -> list[dict]:
    data = json.loads(SEO_PAGES.read_text(encoding="utf-8"))
    return data.get("pages", [])


def research_for(url: str):
    return RESEARCH.get(url, ("Needs research", None, None, "Run seo-keyword-research; update keyword map then this planner"))


def main():
    pages = load_seo_pages()
    solutions = load_solutions()
    by_url = {p["url"]: p for p in pages}

    wb = Workbook()

    # Catalog hubs that own spokes (NOT listed on Core pages)
    hubs_def = [
        ("/solutions", "Hub", "—", "brilliant directories solutions"),
        ("/solutions/seo", "Hub (category)", "/solutions", "brilliant directories schema solutions"),
        ("/solutions/lead-gen", "Hub (category)", "/solutions", "directory lead generation"),
        ("/solutions/member-profiles", "Hub (category)", "/solutions", "directory member profiles"),
        ("/solutions/search", "Hub (category)", "/solutions", "directory search filters"),
        ("/solutions/page-design", "Hub (category)", "/solutions", "directory page design"),
        ("/solutions/content", "Hub (category)", "/solutions", "directory content solutions"),
        ("/solutions/integrations", "Hub (category)", "/solutions", "brilliant directories integrations"),
        ("/solutions/member-management", "Hub (category)", "/solutions", "directory member management"),
        ("/grow-with-ai", "Hub (AI cluster)", "—", "brilliant directories ai"),
        ("/tools", "Hub", "—", "brilliant directories tools"),
        ("/themes", "Hub", "/tools", "brilliant directories themes"),
        ("/blog", "Hub", "—", "brilliant directories blog"),
    ]
    hub_urls = {u for u, *_ in hubs_def}

    # Spokes of those hubs (solutions + AI/tool leaves) — NOT on Core pages
    spoke_pages = [
        ("AI Development", "/ai-development", "/grow-with-ai", "ai brilliant directories development"),
        ("AI for Your Site", "/ai-for-your-site", "/grow-with-ai", "ai for directory website"),
        ("AI Dev Fleet", "/ai-dev-fleet", "/grow-with-ai", "ai development fleet"),
        ("BD Automation", "/bd-automation", "/grow-with-ai", "bd automation"),
        ("BrilliantChat", "/tools/brilliantchat", "/tools", "directory ai chatbot"),
    ]
    spoke_page_urls = {u for _, u, *_ in spoke_pages}

    # ------------------------------------------------------------------ Core pages (standalone — not hubs, not spokes)
    ws0 = wb.active
    ws0.title = "Core pages"
    core_headers = [
        "URL", "Page Title", "Build Status", "Funnel",
        "Primary Keyword", "Secondary Keywords", "Keyword Research Status", "Sprint", "Notes",
    ]
    # Explicit core set from workspace (seo-pages money/trust pages outside hub/spoke trees)
    core_urls = [
        "/",
        "/services",
        "/setup",
        "/hire-developer",
        "/customization",
        "/maintenance",
        "/consultation",
        "/founders-track",
        "/zoom-clinics",
        "/seo-growth",
        "/reviews",
        "/webinars",
        "/case-studies",
        "/about",
        "/contact",
    ]
    core_rows = []
    for url in core_urls:
        assert url not in hub_urls and url not in spoke_page_urls
        p = by_url.get(url, {})
        status = (p.get("status") or "planned").capitalize()
        rstat, primary, secondary, notes = research_for(url)
        core_rows.append((
            url,
            p.get("title") or url,
            status,
            p.get("funnel") or "—",
            primary or p.get("primary_keyword") or "—",
            secondary or "—",
            rstat,
            p.get("sprint") or "—",
            notes,
        ))

    write_rows(ws0, core_headers, core_rows, status_col=3, research_col=7)
    autosize(ws0, [22, 52, 12, 14, 40, 48, 28, 8, 48])
    ws0.row_dimensions[1].height = 30
    for r in range(2, len(core_rows) + 2):
        ws0.row_dimensions[r].height = 40

    # ------------------------------------------------------------------ Hubs
    ws1 = wb.create_sheet("Hubs")
    hub_headers = [
        "Hub Name", "URL", "Build Status", "Funnel", "Parent",
        "Primary Keyword", "Secondary Keywords", "Question Targets (FAQ / AEO)",
        "Do NOT Primary-Target", "Keyword Research Status", "Strategy Notes", "Sprint",
    ]

    hub_rows = [
        (
            "Solutions (master)", "/solutions", "Live", "consideration", "—",
            "brilliant directories solutions", "—", "—",
            "Individual product queries → /solutions/[slug]", "Needs research",
            f"Catalog hub. {len(solutions)} live solution spokes under category hubs.", 2,
        ),
        (
            "SEO & Schema", "/solutions/seo", "Live", "consideration", "/solutions",
            "brilliant directories schema solutions",
            "SEO & Schema Solutions for Brilliant Directories; structured data; JSON-LD; Essential SEO Audit; Strategic SEO Audit",
            "How do I add schema markup to BD?; Does BD include schema by default?; What schema solutions do you offer?; Do I need a schema plugin if BD already has Schema & SEO?; What is an Essential SEO Audit?",
            "brilliant directories schema (platform docs); brilliant directories seo (→ /seo-growth); product long-tails (→ tool pages)",
            "Researched",
            "Own partner catalog phrase. Update seo-pages.json primary from 'brilliant directories schema'.", 2,
        ),
        (
            "Lead Gen", "/solutions/lead-gen", "Live", "consideration", "/solutions",
            "directory lead generation", "—", "—", "—", "Needs research",
            "Category hub — research pass pending.", 2,
        ),
        (
            "Member Profiles", "/solutions/member-profiles", "Live", "consideration", "/solutions",
            "directory member profiles", "—", "—", "—", "Needs research",
            "Category hub — research pass pending.", 2,
        ),
        (
            "Search & Discovery", "/solutions/search", "Live", "consideration", "/solutions",
            "directory search filters", "—", "—", "—", "Needs research",
            "Category hub — research pass pending.", 2,
        ),
        (
            "Page Design", "/solutions/page-design", "Live", "consideration", "/solutions",
            "directory page design", "—", "—", "—", "Needs research",
            "Category hub — research pass pending.", 2,
        ),
        (
            "Content & Engagement", "/solutions/content", "Live", "consideration", "/solutions",
            "directory content solutions", "—", "—", "—", "Needs research",
            "Category hub — research pass pending.", 2,
        ),
        (
            "Integrations", "/solutions/integrations", "Live", "consideration", "/solutions",
            "brilliant directories integrations", "—", "—", "—", "Needs research",
            "Category hub — research pass pending.", 2,
        ),
        (
            "Member Management", "/solutions/member-management", "Live", "consideration", "/solutions",
            "directory member management", "—", "—", "—", "Needs research",
            "Category hub — research pass pending.", 2,
        ),
        (
            "Grow with AI", "/grow-with-ai", "Planned", "consideration", "—",
            "brilliant directories ai", "—", "—", "—", "Needs research",
            "AI cluster parent; spokes on Spokes sheet.", 2,
        ),
        (
            "Tools", "/tools", "Planned", "consideration", "—",
            "brilliant directories tools", "—", "—", "—", "Needs research",
            "Tools hub (planned). Solutions live under /solutions today.", 2,
        ),
        (
            "Themes", "/themes", "Planned", "consideration", "/tools",
            "brilliant directories themes", "—", "—", "—", "Needs research",
            "Themes catalog (planned).", 2,
        ),
        (
            "Blog", "/blog", "Planned", "awareness", "—",
            "brilliant directories blog", "—", "PAA routed off money pages",
            "Do not steal money-page primaries", "Needs research",
            "Supporting content hub (planned).", 2,
        ),
    ]
    write_rows(ws1, hub_headers, hub_rows, status_col=3, research_col=10)
    autosize(ws1, [24, 32, 12, 14, 14, 36, 40, 42, 40, 18, 44, 8])
    ws1.row_dimensions[1].height = 32
    for r in range(2, len(hub_rows) + 2):
        ws1.row_dimensions[r].height = 56

    # ------------------------------------------------------------------ Spokes
    ws2 = wb.create_sheet("Spokes")
    spoke_headers = [
        "Spoke Name", "URL", "Parent Hub", "Build Status", "Funnel / Type",
        "Primary Keyword", "Secondary Keywords", "Question Targets",
        "Keyword Research Status", "Strategy Notes", "Sprint",
    ]

    spoke_rows = []

    for name, url, parent, primary in spoke_pages:
        p = by_url.get(url, {})
        status = (p.get("status") or "planned").capitalize()
        rstat, _p, secondary, notes = research_for(url)
        spoke_rows.append((
            name, url, parent, status, p.get("funnel") or "—",
            primary, secondary or "—", "—", rstat, notes, p.get("sprint") or "—",
        ))

    product_kw_hints = {
        "schema-for-home-page": "schema for home page",
        "essential-seo-audit": "essential seo audit brilliant directories",
        "strategic-seo-audit-advanced": "strategic seo audit brilliant directories",
        "blog-schema-for-blogs-results-page": "blog schema",
        "event-schema-for-event-detail-pages": "event schema",
        "job-schema-for-job-detail-page": "job schema",
        "job-schema-for-jobs-results-page": "job schema (results)",
        "schema-for-members-profile-detail-page": "member profile schema",
    }

    for sol in solutions:
        cat = sol.get("category") or ""
        cat_name, parent_url, _hub_kw = CAT_META.get(cat, CAT_META[""])
        if cat:
            parent = parent_url
            type_label = f"Solution · {cat_name}"
        else:
            parent = "/solutions"
            type_label = "Solution · uncategorized"

        slug = sol["slug"]
        url = f"/solutions/{slug}"
        primary = product_kw_hints.get(slug, "— (product long-tail — needs research)")
        if cat == "seo":
            rstat = "Needs research — product long-tails"
            notes = "Parent hub /solutions/seo is Researched. This leaf owns Marketplace-style long-tail."
        elif not cat:
            rstat = "Needs research"
            notes = "No category in hub map — assign to a /solutions/{category} hub."
        else:
            rstat = "Needs research — product long-tails"
            notes = f"Parent hub {parent_url} still needs category research pass."

        spoke_rows.append((
            sol["title"],
            url,
            parent,
            "Live" if sol.get("status") == "published" else sol.get("status", "Live"),
            type_label,
            primary,
            "structured data; JSON-LD; done-for-you" if cat == "seo" else "—",
            "Product-specific FAQs on detail page",
            rstat,
            notes,
            2,
        ))

    write_rows(ws2, spoke_headers, spoke_rows, status_col=4, research_col=9)
    autosize(ws2, [48, 64, 32, 12, 28, 40, 36, 32, 28, 44, 8])
    ws2.row_dimensions[1].height = 30
    for r in range(2, len(spoke_rows) + 2):
        ws2.row_dimensions[r].height = 32

    # ------------------------------------------------------------------ Research Queue
    ws3 = wb.create_sheet("Research Queue")
    rq_headers = ["Priority", "URL", "Sheet", "Build Status", "Research Status", "Why / Gap", "Action"]
    rq_rows = [
        ("1 — Lock now", "/reviews", "Core pages", "Live", "Researched", "Keyword map cluster done", "Yakin review → lock"),
        ("1 — Lock now", "/zoom-clinics", "Core pages", "Live", "Researched", "Keyword map cluster done", "Yakin review → lock; registry primary = zoom clinics"),
        ("1 — Lock now", "/solutions/seo", "Hubs", "Live", "Researched", "Registry primary still 'brilliant directories schema'", "Yakin review → lock; fix seo-pages.json"),
        ("2 — Confirm", "/", "Core pages", "Live", "Partial — confirm", "Autosuggest + FAQs incomplete", "Run seo-keyword-research"),
        ("2 — Confirm", "/setup", "Core pages", "Planned", "Partial — confirm", "Cost question only", "Research before build"),
        ("2 — Confirm", "/hire-developer", "Core pages", "Planned", "Partial — confirm", "Thin keywords", "Research before build"),
        ("2 — Confirm", "/consultation", "Core pages", "Planned", "Partial — confirm", "Shared KW with /", "Confirm intent split"),
        ("3 — Live no research", "/services", "Core pages", "Live", "Needs research", "Live without brief", "Research"),
        ("3 — Live no research", "/solutions", "Hubs", "Live", "Needs research", "Master hub", "Research"),
        ("3 — Live no research", "/customization", "Core pages", "Live", "Needs research", "Live money page", "Research"),
        ("3 — Live no research", "/webinars", "Core pages", "Live", "Needs research", "Keep ≠ Zoom Clinics", "Research"),
    ]
    for _slug, (name, url, _kw) in [(k, v) for k, v in CAT_META.items() if k and k != "seo"]:
        rq_rows.append((
            "3 — Live category hubs",
            url,
            "Hubs",
            "Live",
            "Needs research",
            f"{name} — sibling of researched /solutions/seo",
            "Run keyword research (same method as seo hub)",
        ))
    rq_rows.extend([
        ("4 — Solutions leaves", f"/solutions/[slug] × {len(solutions)}", "Spokes", "Live", "Needs research — product long-tails", "Each owns Marketplace-style long-tail", "Batch research starting with /solutions/seo children"),
        ("5 — Planned", "/seo-growth", "Core pages", "Planned", "Needs research", "Owns 'brilliant directories seo'", "Research then build"),
        ("5 — Planned", "/grow-with-ai + AI spokes", "Hubs / Spokes", "Planned", "Needs research", "AI cluster", "Research then build"),
        ("5 — Planned", "/about, /contact, /case-studies", "Core pages", "Planned", "Needs research", "In seo-pages.json", "Research then build"),
        ("5 — Planned", "/themes, /blog, /tools", "Hubs", "Planned", "Needs research", "Catalog hubs", "Research then build"),
    ])
    write_rows(ws3, rq_headers, rq_rows, status_col=4, research_col=5)
    autosize(ws3, [22, 40, 14, 12, 28, 44, 36])
    for r in range(2, len(rq_rows) + 2):
        ws3.row_dimensions[r].height = 36

    # ------------------------------------------------------------------ Legend
    ws4 = wb.create_sheet("Legend")
    ws4["A1"] = "BD Growth Suite — Keyword & Hub/Spoke Strategy Planner"
    ws4["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws4.merge_cells("A1:B1")
    legend = [
        ("", ""),
        ("Purpose", "Planning discussion sheet for workspace SEO state. Not a lock source."),
        ("Sources (workspace only)", "data/seo-pages.json + docs/seo-keyword-map.md (research status) + live BD solutions DB/seed (49 products)"),
        ("Do not edit from this sheet", "docs/seo-keyword-map.md stays authoritative for researched clusters until lock"),
        ("", ""),
        ("Sheet", "Contents"),
        ("Core pages", "Standalone site pages that are NOT hubs and NOT spokes (home, services, money/trust pages)"),
        ("Hubs", "Catalog hubs that own spokes (/solutions + categories, AI/tools/themes/blog)"),
        ("Spokes", "Children of hubs — AI/tool leaves + EVERY live /solutions/{slug}"),
        ("Research Queue", "What to lock / confirm / research next"),
        ("", ""),
        ("Research status", "Researched = in keyword-map SERP pass. Partial — confirm = started. Needs research = not done yet."),
        ("Build status", "Live/Planned from seo-pages.json; solutions = published in DB"),
        ("Uncategorized solutions", "3 solutions have empty hub in solutions_hub_map.php — parent = /solutions until assigned"),
    ]
    for i, (a, b) in enumerate(legend, 3):
        ws4.cell(i, 1, a).font = Font(bold=True)
        ws4.cell(i, 2, b).alignment = wrap
    autosize(ws4, [28, 100])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(OUT)
        saved = OUT
    except PermissionError:
        alt = OUT.with_name(OUT.stem + "-new" + OUT.suffix)
        wb.save(alt)
        saved = alt
        print(f"NOTE: {OUT.name} is locked (close Excel). Wrote {alt.name} instead.")
    print(f"Saved: {saved}")
    print(f"Sheets: {wb.sheetnames}")
    print(f"Core pages {len(core_rows)} | Hubs {len(hub_rows)} | Spokes {len(spoke_rows)} (solutions={len(solutions)}) | Queue {len(rq_rows)}")


if __name__ == "__main__":
    main()
